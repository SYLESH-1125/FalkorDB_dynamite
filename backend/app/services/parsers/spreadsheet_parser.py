import csv
import io

import openpyxl


def parse_spreadsheet(file_bytes: bytes, filename: str) -> dict:
    if filename.lower().endswith(".csv"):
        text = file_bytes.decode("utf-8", errors="ignore")
        rows = list(csv.reader(io.StringIO(text)))
        return {
            "page_count": 1,
            "text": text,
            "tables": [{"page": 1, "rows": rows}],
            "used_ocr": False,
        }

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    tables = []
    text_parts = []
    for i, sheet_name in enumerate(wb.sheetnames, start=1):
        sheet = wb[sheet_name]
        rows = [[str(c) if c is not None else "" for c in row] for row in sheet.iter_rows(values_only=True)]
        tables.append({"page": i, "rows": rows})
        text_parts.append("\n".join(" ".join(r) for r in rows))

    return {
        "page_count": len(wb.sheetnames),
        "text": "\n".join(text_parts),
        "tables": tables,
        "used_ocr": False,
    }
