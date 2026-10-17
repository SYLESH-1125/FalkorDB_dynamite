from io import BytesIO

from docx import Document as DocxDocument
from docx.shared import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as PdfImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from app.infrastructure.storage.s3_client import get_object_bytes

ASSET_PREFIX = "/api/reports/assets/"


def _node_text(node: dict) -> str:
    return "".join(c.get("text", "") for c in node.get("content", []) if c.get("type") == "text")


def _table_rows(node: dict) -> list[list[str]]:
    rows = []
    for row in node.get("content", []):
        cells = []
        for cell in row.get("content", []):
            text = " ".join(_node_text(p) for p in cell.get("content", []) if p.get("type") == "paragraph")
            cells.append(text)
        rows.append(cells)
    return rows


def _asset_bytes(src: str) -> bytes | None:
    if not src.startswith(ASSET_PREFIX):
        return None
    key = src[len(ASSET_PREFIX):]
    try:
        return get_object_bytes(key)
    except Exception:
        return None


def render_pdf(title: str, tiptap_doc: dict) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=20, spaceAfter=14, textColor="#171717")
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=14, spaceBefore=16, spaceAfter=8, textColor="#12304a")
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=10.5, leading=15, spaceAfter=8)

    story = []
    for node in tiptap_doc.get("content", []):
        node_type = node.get("type")
        if node_type == "heading":
            level = node.get("attrs", {}).get("level", 2)
            story.append(Paragraph(_node_text(node), h1 if level == 1 else h2))
        elif node_type == "paragraph":
            text = _node_text(node)
            if text:
                story.append(Paragraph(text, body))
        elif node_type == "table":
            rows = _table_rows(node)
            if rows:
                t = Table(rows, hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), "#f2f3f5"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.5, "#e2e5e8"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                ]))
                story.append(t)
                story.append(Spacer(1, 10))
        elif node_type == "image":
            src = node.get("attrs", {}).get("src", "")
            img_bytes = _asset_bytes(src)
            if img_bytes:
                story.append(PdfImage(BytesIO(img_bytes), width=5.5 * inch, height=3.1 * inch))
                story.append(Spacer(1, 10))

    doc.build(story)
    return buf.getvalue()


def render_docx(title: str, tiptap_doc: dict) -> bytes:
    doc = DocxDocument()
    for node in tiptap_doc.get("content", []):
        node_type = node.get("type")
        if node_type == "heading":
            level = node.get("attrs", {}).get("level", 2)
            doc.add_heading(_node_text(node), level=1 if level == 1 else 2)
        elif node_type == "paragraph":
            text = _node_text(node)
            if text:
                doc.add_paragraph(text)
        elif node_type == "table":
            rows = _table_rows(node)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = value
        elif node_type == "image":
            src = node.get("attrs", {}).get("src", "")
            img_bytes = _asset_bytes(src)
            if img_bytes:
                doc.add_picture(BytesIO(img_bytes), width=Inches(6))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
